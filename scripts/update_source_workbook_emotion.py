import os
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(os.environ.get("MYF_SOURCE_XLSX", "MYFdatabase.xlsx")).expanduser().resolve()
TARGET_COLUMNS = {
    "Responses": {"emotion"},
    "121125-3726": {"Emotion"},
}


def replace_value(value):
    if value in {"Amusement", "amusement"}:
        return "Happiness"
    if isinstance(value, str):
        return value.replace("Amusement", "Happiness")
    return value


def main():
    workbook = load_workbook(SOURCE)
    replacements = 0
    for sheet in workbook.worksheets:
        headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
        for header in TARGET_COLUMNS.get(sheet.title, set()):
            column = headers.get(header)
            if not column:
                continue
            for row_index in range(2, sheet.max_row + 1):
                cell = sheet.cell(row=row_index, column=column)
                new_value = replace_value(cell.value)
                if new_value != cell.value:
                    cell.value = new_value
                    replacements += 1
    workbook.save(SOURCE)
    print(f"Updated {SOURCE} with {replacements} replacements")


if __name__ == "__main__":
    main()
