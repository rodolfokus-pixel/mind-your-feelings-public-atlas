import os
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


SOURCE_XLSX = Path(os.environ.get("MYF_SOURCE_XLSX", "MYFdatabase.xlsx")).expanduser().resolve()
NEW_RESPONSES_CSV = Path(os.environ.get("MYF_NEW_RESPONSES_CSV", "Responses-Table 1.csv")).expanduser().resolve()


def clean_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, str):
        return value.strip()
    return value


def canonical_emotion(value):
    text = clean_value(value)
    if isinstance(text, str) and text.lower() in {"amusement", "happy", "happiness", "felicidad"}:
        return "Happiness"
    return text


def main():
    new_rows = pd.read_csv(NEW_RESPONSES_CSV, dtype=str)
    new_rows = new_rows[[column for column in new_rows.columns if column and not str(column).startswith("Unnamed:")]]

    workbook = load_workbook(SOURCE_XLSX)
    sheet = workbook["Responses"]
    headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
    existing_ids = {
        str(sheet.cell(row, 1).value).strip()
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row, 1).value
    }

    appended = 0
    skipped = 0
    for _, source_row in new_rows.iterrows():
        submission_id = clean_value(source_row.get("submission_id"))
        if not submission_id or submission_id in existing_ids:
            skipped += 1
            continue

        output_row = []
        for header in headers:
            if not header:
                output_row.append(None)
                continue
            value = clean_value(source_row.get(header))
            if header == "emotion":
                value = canonical_emotion(value)
            output_row.append(value)

        sheet.append(output_row)
        existing_ids.add(submission_id)
        appended += 1

    workbook.save(SOURCE_XLSX)
    print(f"Appended {appended} new response rows; skipped {skipped} existing or blank rows.")


if __name__ == "__main__":
    main()
